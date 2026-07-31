"""
tests/test_features/test_user_filter_render.py

FEAT-6 — заглушки скрытых участников в готовых документах.
Проверяется главное: текст скрытого участника не попадает в файл.
"""

import json

import pytest

from core.database import DBManager
from features.export.filters import UserFilter
from features.export.generator import (
    DocxGenerator,
    HtmlGenerator,
    JsonGenerator,
    MarkdownGenerator,
)

HIDDEN_TEXT = "СЕКРЕТНЫЙ_ТЕКСТ_БОБА"
VISIBLE_A = "видимый текст алисы"
VISIBLE_C = "видимый текст кэрол"


@pytest.fixture
def db_three_senders():
    """Alice(111) → Bob(222, будет скрыт) → Carol(333, отвечает Бобу)."""
    with DBManager(":memory:") as db:
        db.insert_messages_batch([
            {
                "chat_id": -1001, "message_id": 1,
                "date": "2024-01-15 10:01:00",
                "topic_id": None, "user_id": 111, "username": "Alice",
                "text": VISIBLE_A,
                "media_path": None, "file_type": None, "file_size": None,
                "reply_to_msg_id": None, "post_id": None,
                "is_comment": 0, "from_linked_group": 0,
            },
            {
                "chat_id": -1001, "message_id": 2,
                "date": "2024-01-15 10:02:00",
                "topic_id": None, "user_id": 222, "username": "Bob",
                "text": HIDDEN_TEXT,
                "media_path": None, "file_type": None, "file_size": None,
                "reply_to_msg_id": None, "post_id": None,
                "is_comment": 0, "from_linked_group": 0,
            },
            {
                "chat_id": -1001, "message_id": 3,
                "date": "2024-01-15 10:03:00",
                "topic_id": None, "user_id": 333, "username": "Carol",
                "text": VISIBLE_C,
                "media_path": None, "file_type": None, "file_size": None,
                "reply_to_msg_id": 2, "post_id": None,
                "is_comment": 0, "from_linked_group": 0,
            },
        ])
        yield db


@pytest.fixture
def exclude_bob():
    return UserFilter.make("exclude", [222])


def _read(path):
    with open(path, encoding="utf-8") as fh:
        return fh.read()


# ──────────────────────────────────────────────────────────────────────────────
# Markdown
# ──────────────────────────────────────────────────────────────────────────────

class TestMarkdownPlaceholder:
    def test_hidden_text_absent(self, db_three_senders, exclude_bob, tmp_path):
        gen = MarkdownGenerator(db_three_senders, output_dir=str(tmp_path),
                                user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert HIDDEN_TEXT not in out

    def test_placeholder_present(self, db_three_senders, exclude_bob, tmp_path):
        gen = MarkdownGenerator(db_three_senders, output_dir=str(tmp_path),
                                user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert "Сообщение скрыто" in out

    def test_other_users_untouched(self, db_three_senders, exclude_bob, tmp_path):
        gen = MarkdownGenerator(db_three_senders, output_dir=str(tmp_path),
                                user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert VISIBLE_A in out
        assert VISIBLE_C in out

    def test_author_name_kept(self, db_three_senders, exclude_bob, tmp_path):
        """PLACEHOLDER_SHOW_AUTHOR = True — имя в заглушке остаётся."""
        gen = MarkdownGenerator(db_three_senders, output_dir=str(tmp_path),
                                user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert "Bob" in out

    def test_reply_context_survives(self, db_three_senders, exclude_bob, tmp_path):
        """Ради этого и нужна заглушка: ответ Кэрол не повисает."""
        gen = MarkdownGenerator(db_three_senders, output_dir=str(tmp_path),
                                user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert "в ответ на: Bob" in out

    def test_without_filter_text_present(self, db_three_senders, tmp_path):
        """Контроль: без фильтра текст на месте."""
        gen = MarkdownGenerator(db_three_senders, output_dir=str(tmp_path))
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert HIDDEN_TEXT in out


# ──────────────────────────────────────────────────────────────────────────────
# JSON
# ──────────────────────────────────────────────────────────────────────────────

class TestJsonPlaceholder:
    def test_hidden_text_absent(self, db_three_senders, exclude_bob, tmp_path):
        gen = JsonGenerator(db_three_senders, output_dir=str(tmp_path),
                            user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert HIDDEN_TEXT not in out

    def test_message_count_unchanged(self, db_three_senders, exclude_bob, tmp_path):
        """Скрытое сообщение остаётся в выгрузке — не удаляется."""
        gen = JsonGenerator(db_three_senders, output_dir=str(tmp_path),
                            user_filter=exclude_bob)
        raw = json.loads(_read(gen.generate(-1001, "Тест", period_label="all")[0]))
        records = raw if isinstance(raw, list) else raw.get("messages", raw)
        assert len(records) == 3

    def test_without_filter_text_present(self, db_three_senders, tmp_path):
        gen = JsonGenerator(db_three_senders, output_dir=str(tmp_path))
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert HIDDEN_TEXT in out


# ──────────────────────────────────────────────────────────────────────────────
# HTML
# ──────────────────────────────────────────────────────────────────────────────

class TestHtmlPlaceholder:
    def test_hidden_text_absent(self, db_three_senders, exclude_bob, tmp_path):
        gen = HtmlGenerator(db_three_senders, output_dir=str(tmp_path),
                            user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert HIDDEN_TEXT not in out

    def test_placeholder_present(self, db_three_senders, exclude_bob, tmp_path):
        gen = HtmlGenerator(db_three_senders, output_dir=str(tmp_path),
                            user_filter=exclude_bob)
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert "Сообщение скрыто" in out

    def test_without_filter_text_present(self, db_three_senders, tmp_path):
        gen = HtmlGenerator(db_three_senders, output_dir=str(tmp_path))
        out = _read(gen.generate(-1001, "Тест", period_label="all")[0])
        assert HIDDEN_TEXT in out


# ──────────────────────────────────────────────────────────────────────────────
# Медиа и расшифровки
# ──────────────────────────────────────────────────────────────────────────────

class TestHiddenMediaAndStt:
    def test_media_not_embedded(self, tmp_path):
        """Медиа скрытого участника не встраивается, файл на диске цел."""
        media = tmp_path / "secret_photo.jpg"
        media.write_bytes(b"\xff\xd8\xff")

        with DBManager(":memory:") as db:
            db.insert_messages_batch([{
                "chat_id": -1001, "message_id": 1,
                "date": "2024-01-15 10:01:00",
                "topic_id": None, "user_id": 222, "username": "Bob",
                "text": HIDDEN_TEXT,
                "media_path": str(media), "file_type": "photo",
                "file_size": 3,
                "reply_to_msg_id": None, "post_id": None,
                "is_comment": 0, "from_linked_group": 0,
            }])
            gen = MarkdownGenerator(db, output_dir=str(tmp_path),
                                    user_filter=UserFilter.make("exclude", [222]))
            out = _read(gen.generate(-1001, "Тест", period_label="all")[0])

        assert HIDDEN_TEXT not in out
        assert "secret_photo.jpg" not in out
        assert media.exists(), "файл медиа удалять нельзя"

    def test_transcription_not_leaked(self, tmp_path):
        """
        Расшифровка голосового скрытого участника не должна уехать
        в документ мимо заглушки (stt_map чистится в _apply_user_filter).
        """
        with DBManager(":memory:") as db:
            db.insert_messages_batch([{
                "chat_id": -1001, "message_id": 1,
                "date": "2024-01-15 10:01:00",
                "topic_id": None, "user_id": 222, "username": "Bob",
                "text": "",
                "media_path": None, "file_type": "voice", "file_size": None,
                "reply_to_msg_id": None, "post_id": None,
                "is_comment": 0, "from_linked_group": 0,
            }])
            db.insert_transcription(message_id=1, peer_id=-1001, text=HIDDEN_TEXT)

            gen = MarkdownGenerator(db, output_dir=str(tmp_path),
                                    user_filter=UserFilter.make("exclude", [222]))
            out = _read(gen.generate(-1001, "Тест", period_label="all")[0])

        assert HIDDEN_TEXT not in out


# ──────────────────────────────────────────────────────────────────────────────
# Режим «по постам»: фильтр применяется только к комментариям (D-1, D-2)
# ──────────────────────────────────────────────────────────────────────────────

POST_TEXT       = "ТЕКСТ_ПОСТА_КАНАЛА"
COMMENT_MARIA   = "МАРКЕР_КОММЕНТАРИЯ_МАРИИ"
COMMENT_STT     = "МАРКЕР_РАСШИФРОВКИ_МАРИИ"
COMMENT_IVAN    = "МАРКЕР_КОММЕНТАРИЯ_ИВАНА"
CHANNEL_ID      = -1001
MARIA, IVAN     = 111, 222
POST_NAMES      = {MARIA: "Мария", IVAN: "Иван"}


@pytest.fixture
def db_channel(tmp_path):
    """Канал: пост от имени канала + комментарии Марии (с голосовым) и Ивана."""
    voice = tmp_path / "voice_101.ogg"
    voice.write_bytes(b"x")
    base = {
        "chat_id": CHANNEL_ID, "topic_id": None, "media_path": None,
        "file_type": None, "file_size": None, "from_linked_group": 0,
    }
    with DBManager(":memory:") as db:
        db.insert_messages_batch([
            {**base, "message_id": 100, "date": "2024-01-15 10:00:00",
             "user_id": CHANNEL_ID, "username": "Канал", "text": POST_TEXT,
             "reply_to_msg_id": None, "post_id": None, "is_comment": 0},
            {**base, "message_id": 101, "date": "2024-01-15 10:05:00",
             "user_id": MARIA, "username": "Мария", "text": COMMENT_MARIA,
             "media_path": str(voice), "file_type": "voice", "file_size": 1,
             "reply_to_msg_id": 100, "post_id": 100, "is_comment": 1},
            {**base, "message_id": 102, "date": "2024-01-15 10:06:00",
             "user_id": IVAN, "username": "Иван", "text": COMMENT_IVAN,
             "reply_to_msg_id": 100, "post_id": 100, "is_comment": 1},
        ])
        db.insert_transcription(101, CHANNEL_ID, COMMENT_STT)
        yield db


def _by_posts_text(db, out_dir, generator_cls, user_filter, user_id=None):
    """
    Запускает выгрузку «по постам» и возвращает склеенный текст файлов.

    user_id передаётся как его передаёт UI: при «только выбранные» с одним
    отмеченным ExportWorker кладёт туда его id. Без этого тест не проверяет
    реальный путь — а именно там легаси-фильтр вымывал посты.
    """
    gen = generator_cls(db, output_dir=str(out_dir), user_filter=user_filter)
    if generator_cls is DocxGenerator:
        files = gen.generate(chat_id=CHANNEL_ID, chat_title="Канал",
                             split_mode="post", include_comments=True,
                             user_id=user_id, period_label="fullchat")
        import docx
        return "\n".join(p.text for f in files
                          for p in docx.Document(f).paragraphs)
    files = gen.generate_by_posts(chat_id=CHANNEL_ID, chat_title="Канал",
                                  user_id=user_id, include_comments=True,
                                  period_label="fullchat")
    return "\n".join(open(f, encoding="utf-8").read() for f in files)


ALL_FOUR = [DocxGenerator, MarkdownGenerator, JsonGenerator, HtmlGenerator]


class TestByPostsCommentFilter:
    """
    Режим «по постам»: пост печатается всегда целиком, фильтр действует
    только на комментарии.
    """

    @pytest.mark.parametrize("generator_cls", ALL_FOUR)
    def test_exclude_hides_comment_everywhere(self, db_channel, tmp_path,
                                              generator_cls):
        """
        D-1: три вещи утекали по трём разным причинам, поэтому три проверки.
        DOCX брал комментарии вторым запросом мимо фильтра; расшифровка
        держалась на обнулении file_type в _hide_row(), которого здесь не было.
        """
        uf = UserFilter.make("exclude", [MARIA], POST_NAMES)
        body = _by_posts_text(db_channel, tmp_path, generator_cls, uf)

        assert COMMENT_MARIA not in body, "текст скрытого комментария в файле"
        assert "voice_101" not in body, "медиа скрытого комментария в файле"
        assert COMMENT_STT not in body, "расшифровка скрытого комментария в файле"

    @pytest.mark.parametrize("generator_cls", ALL_FOUR)
    def test_exclude_keeps_post_and_others(self, db_channel, tmp_path,
                                           generator_cls):
        """Пост и чужие комментарии остаются нетронутыми."""
        uf = UserFilter.make("exclude", [MARIA], POST_NAMES)
        body = _by_posts_text(db_channel, tmp_path, generator_cls, uf)

        assert POST_TEXT in body, "пост исчез или стал заглушкой"
        assert COMMENT_IVAN in body, "чужой комментарий пропал"

    @pytest.mark.parametrize("generator_cls", ALL_FOUR)
    def test_include_keeps_post_and_selected_only(self, db_channel, tmp_path,
                                                  generator_cls):
        """
        D-2: раньше «только Мария» не оставляло ни одного поста (посты идут
        от имени канала) и падало с EmptyDataError.
        """
        uf = UserFilter.make("include", [MARIA], POST_NAMES)
        body = _by_posts_text(db_channel, tmp_path, generator_cls, uf,
                              user_id=MARIA)

        assert POST_TEXT in body, "пост должен печататься целиком"
        assert COMMENT_MARIA in body, "комментарий выбранного участника пропал"
        assert COMMENT_IVAN not in body, "комментарий невыбранного в файле"

    @pytest.mark.parametrize("generator_cls", ALL_FOUR)
    def test_include_draws_no_placeholders(self, db_channel, tmp_path,
                                           generator_cls):
        """
        В режиме «только выбранные» заглушек нет вовсе: иначе файл поста
        состоял бы из двадцати строк «Сообщение скрыто».
        """
        uf = UserFilter.make("include", [MARIA], POST_NAMES)
        body = _by_posts_text(db_channel, tmp_path, generator_cls, uf,
                              user_id=MARIA)

        assert "скрыто" not in body.lower()

    @pytest.mark.parametrize("generator_cls", ALL_FOUR)
    def test_post_survives_exclusion_of_its_author(self, db_channel, tmp_path,
                                                   generator_cls):
        """
        Р-1 в самом прямом виде: скрыт автор постов, то есть сам канал.
        Пост всё равно печатается целиком — он то, вокруг чего собран файл,
        и заглушка вместо него оставила бы обсуждение без предмета.
        """
        uf = UserFilter.make("exclude", [CHANNEL_ID], {CHANNEL_ID: "Канал"})
        body = _by_posts_text(db_channel, tmp_path, generator_cls, uf)

        assert POST_TEXT in body, "пост исчез при исключении автора"
        assert COMMENT_MARIA in body and COMMENT_IVAN in body,             "комментарии не должны страдать от исключения канала"
